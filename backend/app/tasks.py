"""M3 任务管理 + 自动周报路由。

- 任务 CRUD：建（支持 parent_id 建子任务）、列（树状 + 父完成度汇总）、改、删（级联子任务）。
- 自动周报：聚合本周 done 任务 + focus_session 工时，返回 Markdown 草稿。

为什么这些放一个文件：任务与周报同源（周报聚合任务），共享 task 表；
M4 拆解出的步骤、M7 专注工时，最终都会落到这里被周报消费，M3 是“数据枢纽”。
"""
from fastapi import APIRouter, HTTPException

from . import db, schemas


router = APIRouter(prefix="/api/tasks", tags=["tasks"])


@router.post("", response_model=dict)
def create_task(payload: schemas.TaskCreate):
    """新建任务；带 parent_id 即为子任务（M4 计划清单会批量调用）。"""
    tid = db.insert_task(
        title=payload.title,
        description=payload.description,
        deadline=payload.deadline,
        urgency=payload.urgency,
        importance=payload.importance,
        parent_id=payload.parent_id,
    )
    return {"id": tid, "ok": True}


@router.get("", response_model=list)
def get_tasks():
    """任务树：顶层任务 + 内嵌 children，父完成度由子任务实时汇总。"""
    return db.list_task_tree()


@router.get("/stale", response_model=list)
def stale_tasks(days: int = 7):
    """防空洞任务：久未更新（默认 7 天）且未完成，供前端标灰 + 应用内提醒。"""
    return db.list_stale_tasks(days)


@router.get("/{task_id}", response_model=dict)
def get_task(task_id: int):
    t = db.get_task(task_id)
    if not t:
        raise HTTPException(status_code=404, detail="task not found")
    return t


@router.patch("/{task_id}", response_model=dict)
def update_task(task_id: int, payload: schemas.TaskUpdate):
    """更新任务；改 title/description 同步搜索索引，改子项冒泡父完成度。"""
    data = payload.model_dump(exclude_unset=True)
    if not data:
        raise HTTPException(status_code=400, detail="nothing to update")
    ok = db.update_task(task_id, **data)
    if not ok:
        raise HTTPException(status_code=404, detail="task not found")
    return {"ok": True, "id": task_id}


@router.delete("/{task_id}", response_model=dict)
def delete_task(task_id: int):
    """删除任务；因 task.parent_id 设了 ON DELETE CASCADE，子任务一并删除。"""
    ok = db.delete_task(task_id)
    if not ok:
        raise HTTPException(status_code=404, detail="task not found")
    return {"ok": True, "id": task_id}


# ---------------------------------------------------------------------------
# 自动周报
# ---------------------------------------------------------------------------
reports_router = APIRouter(prefix="/api/reports", tags=["reports"])


@reports_router.get("/weekly", response_model=schemas.ReportOut)
def weekly_report():
    """生成本周周报草稿（Markdown）。前端展示后可编辑/导出。"""
    return db.generate_weekly_report()


@reports_router.get("/monthly", response_model=schemas.ReportOut)
def monthly_report(year: int | None = None, month: int | None = None):
    """生成(默认当前月)月报草稿（Markdown），与周报同口径，落 report 表 type=monthly。

    为什么单独接口而非 /weekly?period=month：周期边界与展示文案差异明显，
    分开更清晰；前端周报/月报是两个并列视图。
    """
    return db.generate_monthly_report(year, month)


@reports_router.get("/export")
def export_report(type: str = "weekly", format: str = "docx"):
    """把报告导出成真实文件：Word(.docx) 或 PDF(.pdf)。

    满足 PRD M3「导出 Word/PDF」约束：Word 复用 python-docx（与 M1 文档栈一致），
    PDF 用 reportlab（纯 Python、离线、内置 STSong-Light 中文字体）。
    type=weekly|monthly，format=docx|pdf；默认重新聚合当期数据渲染（保证导出内容=当前草稿）。
    """
    from io import BytesIO

    from docx import Document as DocxDocument
    from docx.shared import Pt
    from fastapi.responses import StreamingResponse

    if type == "monthly":
        rep = db.generate_monthly_report()
        title = f"月报（{rep['period_start']} ~ {rep['period_end']}）"
    else:
        rep = db.generate_weekly_report()
        title = f"周报（{rep['period_start']} ~ {rep['period_end']}）"

    if format == "pdf":
        buf = _export_pdf(rep, title)
        fname = f"{'monthly' if type == 'monthly' else 'weekly'}_report_{rep['period_start']}.pdf"
        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={fname}"},
        )

    doc = DocxDocument()
    for line in rep["markdown"].split("\n"):
        if line.startswith("## "):
            h = doc.add_heading(line[3:], level=2)
            h.runs[0].font.size = Pt(14)
        elif line.startswith("# "):
            h = doc.add_heading(line[2:], level=1)
            h.runs[0].font.size = Pt(18)
        elif line.startswith("- "):
            _add_bold_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif line.strip() == "":
            continue
        else:
            _add_bold_runs(doc.add_paragraph(), line)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    fname = f"{'monthly' if type == 'monthly' else 'weekly'}_report_{rep['period_start']}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


def _add_bold_runs(p, text: str) -> None:
    """把 `**加粗**` 片段在段落里渲染成粗体 run（与前端 inline 处理对齐）。"""
    import re
    parts = re.split(r"\*\*(.+?)\*\*", text)
    for i, seg in enumerate(parts):
        if not seg:
            continue
        run = p.add_run(seg)
        if i % 2 == 1:
            run.bold = True


def _export_pdf(rep: dict, title: str) -> "BytesIO":
    """把报告 Markdown 渲染成真实 PDF（reportlab + 中文 CID 字体，离线可用）。

    为什么用 reportlab 而非 docx→pdf 转档：转档依赖本机装 LibreOffice/MS Word，
    离线环境不一定有；reportlab 纯 Python + 自带 STSong-Light 中文字体，零外部依赖，
    符合项目“本地优先/离线可用”的非功能约束。Markdown 仅用到 标题/列表/加粗，逐一映射即可。
    """
    import re
    import html as _html
    from io import BytesIO

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    # 中文需要 CID 字体；STSong-Light 是 Adobe 自带，无需外部字体文件
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    FONT = "STSong-Light"

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=20 * mm, bottomMargin=20 * mm,
        leftMargin=20 * mm, rightMargin=20 * mm,
        title=title,
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontName=FONT, fontSize=18, spaceAfter=8, spaceBefore=4)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontName=FONT, fontSize=14, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName=FONT, fontSize=11, leading=18, spaceAfter=2)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=14, bulletIndent=4, spaceAfter=2)

    def md_to_para(text: str) -> str:
        # 先转义 XML 特殊字符，再把 **加粗** 换成 <b>；** 经转义不受影响
        esc = _html.escape(text)
        return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", esc)

    story = []
    for line in rep["markdown"].split("\n"):
        if not line.strip():
            story.append(Spacer(1, 4))
            continue
        if line.startswith("## "):
            story.append(Paragraph(md_to_para(line[3:]), h2))
        elif line.startswith("# "):
            story.append(Paragraph(md_to_para(line[2:]), h1))
        elif line.startswith("- "):
            story.append(Paragraph(md_to_para(line[2:]), bullet, bulletText="•"))
        else:
            story.append(Paragraph(md_to_para(line), body))
    doc.build(story)
    buf.seek(0)
    return buf

