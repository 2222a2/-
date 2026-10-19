"""M0 资料聚合底座：统一入库流水线 + 四个入库通道的 HTTP 接口。

做了什么（为什么做、逻辑、作用）：
- 平台要“零跳转、不用找资料”，前提是资料能方便地存进来、并随时能被检索 / 被 AI 引用。
- 这里把“存资料”收敛成**一条流水线** ingest_card()：
      写 SQLite 卡片表  →  同步进 FTS5 全文索引  →  （可选）向量化进 chromadb
  四个来源（网页剪藏 / 文件上传 / 任务文档 / 笔记）都走同一条流水线，
  保证“入库即可被搜到、可被 RAG 引用”，不会出现“存了搜不到”的割裂。
- 暴露 /api/knowledge/* 给前端：剪藏、笔记、上传、列表、站内搜索、删除，
  以及笔记系统扩展的「包 CRUD」「结构化笔记（富文本）」「筛选/分组列表」。

与项目的对齐约定（详见 docs/笔记资料系统-方案对齐与落地.md）：
- 接口沿用 /api/knowledge 前缀，不另起 /api/notes。
- content 永远是 FTS 真源（纯文本）；rich_body 存 HTML（展示用），入库前 nh3 净化。
- 文件指针复用 source_id，不引入 file_ref；新增 url 列供网站型。
- 时间区间按 UTC 计算，匹配 SQLite datetime('now') 的 UTC 存储。
- 包路由（/packages）必须注册在 /{card_id} 之前，避免被路径参数吞掉。
"""
from __future__ import annotations

import re
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional

import httpx
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pydantic import BaseModel

from . import config, db
from .ai import rag

router = APIRouter(prefix="/api/knowledge", tags=["knowledge"])

# 上传文件的落盘目录（剪藏/上传的原始文件留档，方便追溯）
UPLOAD_DIR = config.DATA_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

# 富文本净化白名单（后端信任边界，前端净化可被绕过，故必须在后端做）
HTML_ALLOWED_TAGS = {"h1", "h2", "p", "br", "strong", "em", "ul", "ol", "li", "a", "blockquote", "code", "pre"}
HTML_ALLOWED_ATTRS = {"a": {"href"}}


def sanitize_html(html_str: str) -> str:
    """用 nh3 按白名单净化用户 HTML，剥掉 <script>/onclick/iframe 等。失败则退化为纯文本。"""
    try:
        import nh3
        return nh3.clean(html_str, tags=HTML_ALLOWED_TAGS, attributes=HTML_ALLOWED_ATTRS)
    except Exception:
        return re.sub(r"<[^>]+>", "", html_str or "")


def _plain_from_html(html_str: str) -> str:
    """把净化后/未净化的 HTML 抽成纯文本（用于写 content / FTS 索引）。"""
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html_str or "")).strip()


# ---------------------------------------------------------------------------
# 统一入库流水线（底座核心）
# ---------------------------------------------------------------------------
def ingest_card(
    title: str,
    body: str,
    source_type: str,
    source_id: Optional[str] = None,
    url: Optional[str] = None,
    rich_body: Optional[str] = None,
    item_type: Optional[str] = None,
    package_id: Optional[int] = None,
) -> int:
    """一条资料从“任意来源”进入平台的唯一入口。

    逻辑：
    1. 落 SQLite knowledge_card（insert_knowledge_card 内部已同步 FTS5 全文索引）；
       - body 为纯文本（FTS 真源）；
       - rich_body 为 HTML 展示内容（若有），由调用方先 sanitize_html 净化。
    2. 若启用 chromadb（RAG_BACKEND=chromadb 且可用），把正文向量化写进向量库。
    返回新卡片 id。
    """
    cid = db.insert_knowledge_card(
        title=title,
        content=body,
        source_type=source_type,
        source_id=source_id,
        rich_body=rich_body,
        item_type=item_type,
        package_id=package_id,
        url=url,
    )
    rag.index_into_chroma(cid, body, {"title": title, "source": source_type, "url": url or ""})
    return cid


# ---------------------------------------------------------------------------
# 文本抽取工具（来源不同，正文抽取方式不同）
# ---------------------------------------------------------------------------
def extract_from_html(html: str, fallback_title: str = "") -> Dict[str, str]:
    """从网页 HTML 抽正文 + 标题。优先 readability-lxml，缺失时降级为轻量清洗。"""
    title = fallback_title
    try:
        from readability import Document  # 可选依赖
        doc = Document(html)
        body = doc.summary()
        title = doc.short_title() or title
    except Exception:
        html_no_script = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", html, flags=re.S | re.I)
        mt = re.search(r"<title[^>]*>(.*?)</title>", html_no_script, flags=re.S | re.I)
        if mt and not title:
            title = mt.group(1).strip()
        body = re.sub(r"<[^>]+>", " ", html_no_script)
        body = re.sub(r"\s+", " ", body).strip()
    body_text = re.sub(r"<[^>]+>", " ", body)
    body_text = re.sub(r"\s+", " ", body_text).strip()
    return {"title": (title or fallback_title).strip()[:200], "body": body_text}


def extract_from_file(filename: str, raw: bytes) -> str:
    """从上传文件抽正文。txt/md 直接读；pdf/docx 用可选依赖；其它类型尽力而为。"""
    ext = Path(filename).suffix.lower()
    if ext in (".txt", ".md", ".markdown", ".csv", ".json"):
        for enc in ("utf-8", "gb18030", "latin-1"):
            try:
                return raw.decode(enc)
            except Exception:
                continue
        return raw.decode("utf-8", errors="ignore")
    if ext == ".pdf":
        try:
            import pdfplumber  # 可选依赖
            text_parts: List[str] = []
            with pdfplumber.open(__import__("io").BytesIO(raw)) as pdf:
                for page in pdf.pages:
                    text_parts.append(page.extract_text() or "")
            return "\n".join(text_parts)
        except Exception as exc:
            return f"[PDF 抽取失败：未安装 pdfplumber 或文件损坏] {exc}"
    if ext in (".docx", ".doc"):
        try:
            import docx  # python-docx，可选依赖
            d = docx.Document(__import__("io").BytesIO(raw))
            return "\n".join(p.text for p in d.paragraphs)
        except Exception as exc:
            return f"[DOCX 抽取失败：未安装 python-docx 或文件损坏] {exc}"
    try:
        return raw.decode("utf-8", errors="ignore")
    except Exception:
        return f"[不支持的文件类型：{ext}]"


# ---------------------------------------------------------------------------
# 请求模型
# ---------------------------------------------------------------------------
class ClipIn(BaseModel):
    url: str


class NoteIn(BaseModel):
    """笔记请求：兼容旧调用（只传 content 纯文本），也支持富文本结构化笔记。

    - content：纯文本（AI “记一下” / 旧接口）。
    - rich_body：HTML 富文本（前端 TipTap 输出），入库前净化。
    - item_type：note | website | file。
    - url：网站型跳转地址。
    - package_id：归属包（NULL=未归档）。
    """
    title: Optional[str] = None
    content: Optional[str] = None
    rich_body: Optional[str] = None
    item_type: str = "note"
    url: Optional[str] = None
    package_id: Optional[int] = None


class PackageIn(BaseModel):
    name: str
    parent_id: Optional[int] = None
    color: Optional[str] = None
    sort_order: int = 0


class PackagePatch(BaseModel):
    name: Optional[str] = None
    color: Optional[str] = None
    sort_order: Optional[int] = None
    parent_id: Optional[int] = None


class CardPatch(BaseModel):
    title: Optional[str] = None
    content: Optional[str] = None
    rich_body: Optional[str] = None
    item_type: Optional[str] = None
    url: Optional[str] = None
    package_id: Optional[int] = None


# ---------------------------------------------------------------------------
# 工具：时间区间（UTC，匹配 SQLite datetime('now') 存储）
# ---------------------------------------------------------------------------
def time_range_utc(tr: Optional[str]):
    if not tr:
        return None, None
    now = datetime.utcnow()
    if tr == "today":
        start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    elif tr == "week":
        start = now - timedelta(days=now.weekday())
        start = start.replace(hour=0, minute=0, second=0, microsecond=0)
    elif tr == "month":
        start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    else:
        return None, None
    fmt = "%Y-%m-%d %H:%M:%S"
    return start.strftime(fmt), now.strftime(fmt)


def group_by_pinyin(cards: list[dict]) -> list[dict]:
    """按标题拼音首字母归组；组按 A-Z 排序，组内按拼音排。"""
    groups: Dict[str, list] = {}
    for c in cards:
        key = (c.get("title_pinyin") or c.get("title") or "?")
        initial = key[0].upper() if key and key[0].isalpha() else "#"
        groups.setdefault(initial, []).append(c)
    out = []
    for g in sorted(groups.keys()):
        items = sorted(groups[g], key=lambda x: (x.get("title_pinyin") or ""))
        out.append({"group": g, "items": items})
    return out


def group_by_time(cards: list[dict], now: Optional[datetime] = None) -> list[dict]:
    """按今天/本周/本月/更早归组（UTC）。"""
    now = now or datetime.utcnow()
    week_start = (now - timedelta(days=now.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    buckets: Dict[str, list] = {"今天": [], "本周": [], "本月": [], "更早": []}
    for c in cards:
        try:
            d = datetime.strptime(c.get("created_at") or "", "%Y-%m-%d %H:%M:%S")
        except Exception:
            buckets["更早"].append(c)
            continue
        if d.date() == now.date():
            buckets["今天"].append(c)
        elif d >= week_start:
            buckets["本周"].append(c)
        elif d.year == now.year and d.month == now.month:
            buckets["本月"].append(c)
        else:
            buckets["更早"].append(c)
    return [{"group": g, "items": buckets[g]} for g in ("今天", "本周", "本月", "更早") if buckets[g]]


# ---------------------------------------------------------------------------
# 接口
# ---------------------------------------------------------------------------
@router.post("/clip", summary="网页剪藏")
def clip(body: ClipIn):
    """接收一个网页 URL，后端抓取正文 → 抽取 → 走入库流水线（source=web_clip）。"""
    try:
        resp = httpx.get(body.url, timeout=15, follow_redirects=True,
                         headers={"User-Agent": "Mozilla/5.0 WorkLearn-Clip"})
        resp.raise_for_status()
        html = resp.text
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"抓取网页失败：{exc}")

    extracted = extract_from_html(html, fallback_title=body.url)
    cid = ingest_card(
        title=extracted["title"] or body.url,
        body=extracted["body"],
        source_type="web_clip",
        url=body.url,
    )
    return {"id": cid, "title": extracted["title"], "url": body.url, "ok": True}


@router.post("/notes", summary="手动笔记 / 结构化笔记")
def add_note(body: NoteIn):
    """手动写入一条笔记。兼容旧调用（只传 content 纯文本），也支持富文本结构化笔记。

    - 传 rich_body：净化后存 rich_body，content 同步写纯文本（FTS 真源）。
    - 只传 content：按旧逻辑存纯文本，rich_body 留空（展示层回退渲染 content）。
    """
    has_rich = bool(body.rich_body and body.rich_body.strip())
    has_plain = bool(body.content and body.content.strip())
    if not has_rich and not has_plain:
        raise HTTPException(status_code=400, detail="笔记内容不能为空")

    if has_rich:
        clean = sanitize_html(body.rich_body)
        plain = _plain_from_html(clean)
        title = (body.title or plain[:20] or "无标题笔记").strip()
        cid = ingest_card(
            title=title,
            body=plain,
            source_type="note",
            rich_body=clean,
            item_type=body.item_type or "note",
            package_id=body.package_id,
            url=body.url,
        )
    else:
        content = body.content.strip()
        title = (body.title or content[:20]).strip()
        cid = ingest_card(
            title=title,
            body=content,
            source_type="note",
            item_type=body.item_type or "note",
            package_id=body.package_id,
            url=body.url,
        )
    return {"id": cid, "title": title, "ok": True}


@router.post("/uploads", summary="文件上传入库")
async def upload(file: UploadFile = File(...)):
    """上传文件（txt/md/pdf/docx…），抽取正文 → 入库（source=upload）。"""
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="空文件")
    safe_name = re.sub(r"[^\w.\-]", "_", file.filename or "file")
    save_path = UPLOAD_DIR / safe_name
    save_path.write_bytes(raw)
    body = extract_from_file(safe_name, raw)
    title = Path(safe_name).stem[:200]
    cid = ingest_card(
        title=title,
        body=body,
        source_type="upload",
        source_id=safe_name,
        url=str(save_path),
    )
    return {"id": cid, "title": title, "file": safe_name, "ok": True}


@router.get("", summary="列表 / 筛选 / 分组 / 站内搜索")
def list_or_search(
    q: Optional[str] = None,
    source_type: Optional[str] = None,
    item_type: Optional[str] = None,
    package_id: Optional[int] = None,
    include_descendants: int = 0,
    time_range: Optional[str] = None,
    group_by: Optional[str] = None,
    limit: int = 200,
):
    """无 q：列出/筛选/分组资料卡片；有 q：在卡片范围内做全文检索。

    参数：
    - q：FTS5 全文检索（标题+正文）。
    - item_type：按类型筛选（note|website|file）；缺省时回退用 source_type（向后兼容）。
    - package_id + include_descendants=1：筛选某包及其子孙的资料。
    - time_range：today|week|month（UTC 区间）。
    - group_by：none|title|time。
    """
    if q:
        hits = db.search_cards(q, limit=limit)
        return {"mode": "search", "query": q, "results": hits}

    pkg_ids = None
    if package_id is not None:
        pkg_ids = [package_id]
        if include_descendants:
            pkg_ids = list(set(pkg_ids + list(db.get_descendant_package_ids(package_id))))

    start, end = time_range_utc(time_range)
    filter_type = item_type or source_type
    rows = db.query_cards(
        item_type=filter_type, package_ids=pkg_ids,
        time_start=start, time_end=end, limit=limit,
    )

    if group_by == "title":
        return {"mode": "group", "groups": group_by_pinyin(rows)}
    if group_by == "time":
        return {"mode": "group", "groups": group_by_time(rows)}
    return {"mode": "list", "count": len(rows), "cards": rows}


# ---------------------------------------------------------------------------
# 笔记包（note_package）接口 —— 必须注册在 /{card_id} 之前
# ---------------------------------------------------------------------------
@router.get("/packages", summary="包扁平列表")
def list_packages():
    """返回所有包（含 parent_id），前端递归构树。"""
    return {"packages": db.list_packages()}


@router.post("/packages", summary="新建包")
def create_package(body: PackageIn):
    if not body.name.strip():
        raise HTTPException(status_code=400, detail="包名不能为空")
    if body.parent_id is not None and db.get_package(body.parent_id) is None:
        raise HTTPException(status_code=400, detail="父包不存在")
    pid = db.create_package(
        name=body.name.strip(),
        parent_id=body.parent_id,
        color=body.color,
        sort_order=body.sort_order,
    )
    return {"id": pid, "ok": True}


@router.patch("/packages/{pkg_id}", summary="改包名/颜色/移动")
def update_package(pkg_id: int, body: PackagePatch):
    if db.get_package(pkg_id) is None:
        raise HTTPException(status_code=404, detail="包不存在")
    try:
        ok = db.update_package(
            pkg_id, name=body.name, color=body.color,
            sort_order=body.sort_order, parent_id=body.parent_id,
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    if not ok:
        return {"ok": True, "changed": False}
    return {"ok": True, "changed": True}


@router.delete("/packages/{pkg_id}", summary="删包")
def delete_package(pkg_id: int):
    """删包：子资料 package_id 置 NULL 回未归档，子包变顶层。"""
    if db.get_package(pkg_id) is None:
        raise HTTPException(status_code=404, detail="包不存在")
    db.delete_package(pkg_id)
    return {"ok": True, "id": pkg_id}


# ---------------------------------------------------------------------------
# 单条资料（/{card_id}）—— 注册在 /packages 之后
# ---------------------------------------------------------------------------
@router.get("/{card_id}", summary="查看单条")
def get_one(card_id: int):
    card = db.get_knowledge_card(card_id)
    if not card:
        raise HTTPException(status_code=404, detail="卡片不存在")
    return card


@router.patch("/{card_id}", summary="更新单条")
def patch_one(card_id: int, body: CardPatch):
    """更新标题/正文/类型/归属包/链接。rich_body 若提供则净化并重算 content 纯文本。"""
    card = db.get_knowledge_card(card_id)
    if not card:
        raise HTTPException(status_code=404, detail="卡片不存在")

    conn = db.get_conn()
    sets, vals = [], []
    if body.title is not None:
        sets.append("title=?"); vals.append(body.title)
        sets.append("title_pinyin=?"); vals.append(db._pinyin(body.title))
    if body.package_id is not None:
        sets.append("package_id=?"); vals.append(body.package_id)
    if body.item_type is not None:
        sets.append("item_type=?"); vals.append(body.item_type)
    if body.url is not None:
        sets.append("url=?"); vals.append(body.url)
    if body.rich_body is not None:
        clean = sanitize_html(body.rich_body)
        sets.append("rich_body=?"); vals.append(clean)
        sets.append("content=?"); vals.append(_plain_from_html(clean))
    elif body.content is not None:
        sets.append("content=?"); vals.append(body.content)
    if not sets:
        return card
    conn.execute(f"UPDATE knowledge_card SET {','.join(sets)} WHERE id=?", vals + [card_id])
    conn.commit()
    conn.close()
    updated = db.get_knowledge_card(card_id)
    db.upsert_search("knowledge_card", card_id, updated["title"] or "", updated["content"] or "")
    return updated


@router.delete("/{card_id}", summary="删除")
def delete_one(card_id: int):
    """删除卡片，并同步清理全文索引与向量库，避免脏数据。"""
    if not db.delete_knowledge_card(card_id):
        raise HTTPException(status_code=404, detail="卡片不存在")
    rag.remove_from_chroma(card_id)
    return {"ok": True, "id": card_id}
